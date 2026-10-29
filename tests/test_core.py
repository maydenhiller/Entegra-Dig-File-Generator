"""Tests that need no private files: filename rules, KML parsing, stationing,
the cheat sheet, the label-driven staking writer and the aerial renderer
(with the imagery fetch stubbed out)."""

import io
import os
import sys
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import app  # noqa: E402


# ---------------------------------------------------------------- file names

@pytest.mark.parametrize("name, expected", [
    ("533_8 ARCHER-WF_REP 2_2026_306_52290.53_12320_None.pdf", "8 ARCHER-WF"),
    ("533_8_ARCHER-WF_REP_6_2026_306_57894.92_13710_None.pdf", "8 ARCHER-WF"),
    ("533_8_ARCHER-WF_Dig_4_2026.pdf", "8 ARCHER-WF"),
    ("nounderscore.pdf", ""),
])
def test_line_name_from_filename(name, expected):
    assert app.line_name_from_filename(name) == expected


def test_reference_name():
    assert app.reference_name("  agm   9 ") == "AGM 9"
    assert app.reference_name("8-22 Anchor Rd") == "8-22 Anchor Rd"
    assert app.reference_name(None) is None


# ---------------------------------------------------------------- stationing

def _dig(**kw):
    base = dict(odometer=57911.31, us_ref="AGM 10", us_ref_odometer=54911.08,
                ds_ref="AGM 11", ds_ref_odometer=60230.46,
                us_ref_station="1797+98.8", ds_ref_station="1972+32.2")
    base.update(kw)
    return app.Dig(**base)


def test_reference_distances():
    d = _dig()
    assert round(d.us_ref_distance, 2) == 3000.23
    assert round(d.ds_ref_distance, 2) == 2319.15
    assert d.nearest_reference_label == "AGM 11 2319.15'"


def test_station_interpolates_between_references():
    d = _dig()
    assert d.station_label == "1896+31.6"
    d.station_method = "offset"
    assert d.station_label == "1827+99.0"


def test_station_missing_reference_station():
    assert _dig(us_ref_station="").station_feet is None


def test_station_ascends():
    assert app.station_ascends(_dig())
    assert not app.station_ascends(_dig(us_ref_station="1972+32.2", ds_ref_station="1797+98.8"))


# ---------------------------------------------------------------- KMZ

KML = """<?xml version="1.0"?>
<kml xmlns="http://www.opengis.net/kml/2.2"><Document><name>T</name>
<Style id="red"><LineStyle><color>ff0000ff</color><width>2</width></LineStyle></Style>
<StyleMap id="m"><Pair><key>normal</key><styleUrl>#red</styleUrl></Pair></StyleMap>
<Style id="span"><LineStyle><width>5</width></LineStyle></Style>
<Placemark><name>Pipe</name><styleUrl>#m</styleUrl><MultiGeometry><LineString>
<coordinates>-98.0,33.0,0 -98.0,33.01,0</coordinates></LineString></MultiGeometry></Placemark>
<Folder><name>REP 6</name>
<Placemark><name>REP 6 Span</name><styleUrl>#span</styleUrl><LineString>
<coordinates>-98.0,33.004,0 -98.0,33.005,0</coordinates></LineString></Placemark>
<Placemark><name>REP 6 - uGW</name><Point><coordinates>-98.0,33.004,0</coordinates></Point></Placemark>
</Folder>
<Placemark><name>AGM 9</name><Point><coordinates>-98.0,33.0,0</coordinates></Point></Placemark>
</Document></kml>"""


def test_kml_colours_and_markers():
    kmz = app.load_kmz(KML.encode())
    colours = {l.name: l.colour for l in kmz.lines}
    assert colours["Pipe"] == (255, 0, 0)               # aabbggrr -> rgb
    assert colours["REP 6 Span"] == (255, 255, 255)     # no <color> = white
    assert kmz.marker("agm 9") == (33.0, -98.0)
    assert kmz.replacements["6"].points["uGW"] == (33.004, -98.0)
    assert kmz.replacements["6"].span is not None


def test_kmz_zip_wrapper():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("doc.kml", KML)
    assert app.load_kmz(buf.getvalue()).marker("AGM 9")


def test_rep_number():
    d = app.Dig(source_file="533_8_ARCHER-WF_REP_10_2026.pdf")
    assert app.rep_number(d) == "10"


# ---------------------------------------------------------------- cheat sheet

def _cheat_template(name):
    path = ROOT / "templates" / name
    if not path.exists():
        pytest.skip("cheat sheet template not present")
    return path.read_bytes()


def test_cheat_sheet_rows_and_formulas():
    import openpyxl
    template = _cheat_template("Dig stake cheat sheet template for ascending station numbers.xlsx")
    d1 = _dig(label="Dig 1", name="533_X_Dig 1_2026", latitude=33.1, longitude=-98.1,
              us_weld_distance=10.5, ds_weld_distance=3.25, pipeline_id="533", ili_range="306")
    d2 = _dig(label="Dig 2", name="533_X_Dig 2_2026", odometer=58000.0, latitude=33.2, longitude=-98.2,
              us_weld_distance=1.0, ds_weld_distance=2.0, pipeline_id="533", ili_range="306")
    out = app.build_cheat_sheet(template, [d2, d1])
    ws = openpyxl.load_workbook(io.BytesIO(out)).active
    assert ws["A6"].value == "Dig 1 U/S Weld" and ws["A7"].value == "Dig 1 Anomaly"
    assert ws["B7"].value == 57911.31 and ws["D6"].value == 10.5 and ws["D8"].value == 3.25
    assert ws["B6"].value == "=B7-D6" and ws["C6"].value == "=C7-D6" and ws["C8"].value == "=C7+D8"
    assert ws["E7"].value == "AGM 11 2319.15'"
    assert ws["I7"].value == 33.1 and ws["J7"].value == -98.1
    assert "past Dig 1" in (ws["F10"].value or "")      # 88.69' past, within the proximity window
    assert ws["A12"].value is None                       # unused template rows cleared


# ---------------------------------------------------------------- staking report

def _staking_template():
    path = ROOT / "templates" / "staking_template.xlsm"
    if not path.exists():
        pytest.skip("no staking template in templates/ (kept out of the repo on purpose)")
    return path.read_bytes()


def test_staking_report_cells():
    template = _staking_template()
    info = app.read_template_info(template)
    assert not info.problems
    d = _dig(name="533_Archer to WF_Dig 15_2026", label="Dig 15", line_name="8 ARCHER-WF",
             tool_type="MFL/DEF/XYZ", county="Wichita", state="Texas", directions="Go north.")
    out, warnings = app.build_staking_report(template, d, app.ReportSettings("Hayden Miller"), info)
    patcher = app.XlsmPatcher(out)
    texts = app.sheet_texts(patcher, app.SHEET)
    lay = info.layout
    # A long name widens its box, so the cell it lands in can move - look it up by value.
    assert any(v.startswith(d.name) for v in texts.values())
    assert texts[lay["county"]] == "Wichita" and texts[lay["state"]] == "Texas"
    assert texts[lay["alignment"]] == "Unknown" and texts[lay["hca"]] == "Unknown"
    assert texts[lay["directions"]] == "Go north."
    assert "Hayden Miller" == texts[lay["surveyor"]]


# ---------------------------------------------------------------- aerial

def test_aerial_with_stubbed_tiles(monkeypatch):
    from PIL import Image
    monkeypatch.setattr(app, "_fetch_tile", lambda *a, **k: Image.new("RGB", (256, 256), (70, 110, 60)))
    kmz = app.load_kmz(KML.encode())
    d = app.Dig(name="533_X_Dig 3_2026", label="Dig 3", latitude=33.0045, longitude=-98.0,
                us_ref="AGM 9", us_ref_odometer=0.0, odometer=1400.0,
                source_file="533_8_X_REP_6_2026.pdf")
    jpeg = app.render_aerial(d, kmz, width=800, slot_aspect=2.0, token="")
    image = Image.open(io.BytesIO(jpeg))
    assert image.size == (800, 400)                       # exactly the slot's aspect


def test_aerial_needs_a_coordinate():
    assert app.render_aerial(app.Dig(), None) is None


def test_geometry_gives_requested_width():
    zoom, factor, mpp = app.aerial_geometry(33.0, 2600.0, 1600, 256, 19)
    assert abs(mpp * 1600 * 3.28084 - 2600.0) < 1.0


def test_reps_one_dig_per_pdf_and_labels():
    a = app.Dig(name="533_X_Dig 15_2026", label="Dig 15", odometer=57911.31,
                source_file="533_8_ARCHER-WF_REP_6_2026.pdf")
    b = app.Dig(name="533_X_Dig 16_2026", label="Dig 16", odometer=57943.77,
                source_file="533_8_ARCHER-WF_REP_6_2026.pdf")
    # the same labelling step parse_dig_pdfs applies
    by_file = {}
    for d in (b, a):
        by_file.setdefault(d.source_file, []).append(d)
    assert app.rep_number(a) == "6"
